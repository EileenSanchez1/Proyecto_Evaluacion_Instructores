"""
Generación de reportes en Excel (.xlsx) para pendientes de evaluación
e instructores sin foto de perfil.
"""
from __future__ import annotations

import io
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


def _estilo_cabecera(ws, row: int, cols: int):
    fill = PatternFill("solid", fgColor="39A900")
    font = Font(bold=True, color="FFFFFF")
    thin = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin


def _auto_ancho(ws, max_width: int = 45):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = 0
        for cell in col:
            val = str(cell.value) if cell.value is not None else ""
            length = max(length, min(len(val), max_width))
        ws.column_dimensions[letter].width = max(12, length + 2)


def excel_pendientes_evaluacion(data: dict[str, Any]) -> bytes:
    """
    data: respuesta de pendientes_evaluacion
    columnas: Ficha, Aprendiz, Documento, Correo, Instructores pendientes, Total pendientes
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Pendientes evaluación"

    ws["A1"] = "Reporte de aprendices con evaluaciones pendientes"
    ws["A1"].font = Font(bold=True, size=14, color="1F2937")
    ws.merge_cells("A1:F1")
    ws["A2"] = f"Periodo ID: {data.get('periodo_id') or '—'}  |  Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws["A2"].font = Font(color="6B7280", size=10)
    ws.merge_cells("A2:F2")

    headers = [
        "Ficha",
        "Aprendiz",
        "Documento",
        "Correo",
        "Instructores pendientes",
        "Total pendientes",
    ]
    for i, h in enumerate(headers, 1):
        ws.cell(row=4, column=i, value=h)
    _estilo_cabecera(ws, 4, len(headers))

    row = 5
    for p in data.get("pendientes") or []:
        nombres_inst = ", ".join(
            i.get("nombre", "") for i in (p.get("instructores_pendientes") or [])
        )
        ws.cell(row=row, column=1, value=p.get("ficha") or p.get("numero_ficha") or "")
        ws.cell(row=row, column=2, value=p.get("nombre") or "")
        ws.cell(row=row, column=3, value=p.get("documento") or "")
        ws.cell(row=row, column=4, value=p.get("correo") or "")
        ws.cell(row=row, column=5, value=nombres_inst)
        ws.cell(row=row, column=6, value=p.get("total_pendientes") or 0)
        row += 1

    if row == 5:
        ws.cell(row=5, column=1, value="No hay pendientes en el filtro actual.")
        ws.merge_cells("A5:F5")

    _auto_ancho(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def excel_instructores_sin_foto(data: dict[str, Any]) -> bytes:
    """data: respuesta de instructores_sin_foto"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sin foto de perfil"

    ws["A1"] = "Instructores sin foto de perfil"
    ws["A1"].font = Font(bold=True, size=14, color="1F2937")
    ws.merge_cells("A1:E1")
    ws["A2"] = f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}  |  Total: {data.get('total', 0)}"
    ws["A2"].font = Font(color="6B7280", size=10)
    ws.merge_cells("A2:E2")

    headers = ["ID", "Nombre", "Correo institucional", "Teléfono", "Tiene usuario"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=4, column=i, value=h)
    _estilo_cabecera(ws, 4, len(headers))

    row = 5
    for inst in data.get("instructores") or []:
        ws.cell(row=row, column=1, value=inst.get("id_instructor"))
        ws.cell(row=row, column=2, value=inst.get("nombre") or "")
        ws.cell(row=row, column=3, value=inst.get("correo") or "")
        ws.cell(row=row, column=4, value=inst.get("telefono") or "")
        ws.cell(row=row, column=5, value="Sí" if inst.get("tiene_usuario") else "No")
        row += 1

    if row == 5:
        ws.cell(row=5, column=1, value="Todos los instructores tienen foto de perfil.")
        ws.merge_cells("A5:E5")

    _auto_ancho(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def excel_consolidado_evaluacion(data: dict) -> bytes:
    """
    Excel estilo reporte SENA de referencia:
    - Hoja GENERAL con instructor / rol / aprendices / promedio %
    - Una hoja por instructor con conteos Siempre / Algunas veces / Nunca
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "GENERAL"

    titulo = data.get("titulo") or "CONSOLIDADO GENERAL EVALUACIÓN INSTRUCTORES"
    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, size=14, color="1F2937")
    ws.merge_cells("A1:E1")
    ws["A2"] = data.get("periodo_nombre") or ""
    ws["D2"] = f"FECHA APLICACIÓN {data.get('fecha_aplicacion') or '—'}"
    ws["D2"].font = Font(size=10, color="6B7280")

    headers = ["INSTRUCTOR", "ROL", "APRENDICES QUE EVALUARON", "PROMEDIO"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=4, column=i, value=h)
    _estilo_cabecera(ws, 4, len(headers))

    row = 5
    for inst in data.get("instructores") or []:
        ws.cell(row=row, column=1, value=inst.get("nombre") or "")
        ws.cell(row=row, column=2, value=inst.get("rol") or "Técnico")
        ws.cell(row=row, column=3, value=inst.get("aprendices_evaluaron") or 0)
        pct = inst.get("porcentaje")
        ws.cell(row=row, column=4, value=f"{pct} %" if pct is not None else "—")
        row += 1
    if row == 5:
        ws.cell(row=5, column=1, value="Sin datos de evaluación.")
    _auto_ancho(ws)

    fill_sec = PatternFill("solid", fgColor="E8F5E9")
    for inst in data.get("instructores") or []:
        nombre = (inst.get("nombre") or "Instructor").strip()
        safe = "".join(c if c.isalnum() or c in " _-" else "_" for c in nombre)[:28] or "Inst"
        base, n = safe, 1
        while safe in wb.sheetnames:
            safe = f"{base[:25]}_{n}"
            n += 1
        wsi = wb.create_sheet(title=safe)
        wsi["A1"] = data.get("titulo") or "REPORTE EVALUACIÓN INSTRUCTORES"
        wsi["A1"].font = Font(bold=True, size=12)
        wsi["A2"] = data.get("periodo_nombre") or ""
        wsi["C2"] = f"FECHA APLICACIÓN {data.get('fecha_aplicacion') or '—'}"
        wsi["A3"] = f"INSTRUCTOR {nombre}"
        wsi["A3"].font = Font(bold=True)
        wsi["C3"] = f"ROL {inst.get('rol') or 'Técnico'}"

        r = 5
        for etiqueta, key in [("Siempre", "siempre"), ("Algunas Veces", "algunas"), ("Nunca", "nunca")]:
            wsi.cell(row=r, column=1, value="Pregunta").font = Font(bold=True)
            wsi.cell(row=r, column=2, value=etiqueta).font = Font(bold=True)
            wsi.cell(row=r, column=1).fill = fill_sec
            wsi.cell(row=r, column=2).fill = fill_sec
            r += 1
            for p in inst.get("preguntas") or []:
                orden = p.get("orden")
                enunciado = p.get("enunciado") or ""
                label = f"{orden}.{enunciado}" if orden else enunciado
                wsi.cell(row=r, column=1, value=label)
                wsi.cell(row=r, column=2, value=p.get(key) or 0)
                r += 1
            r += 1

        r += 1
        wsi.cell(row=r, column=1, value="CONSOLIDADO EVALUACIÓN").font = Font(bold=True, size=11)
        r += 1
        cons = ["Total Aprendices", "Total Preguntas", "Total Respuestas", "Siempre", "Algunas Veces", "Nunca", "Promedio Instructor"]
        for i, h in enumerate(cons, 1):
            wsi.cell(row=r, column=i, value=h)
        _estilo_cabecera(wsi, r, len(cons))
        r += 1
        wsi.cell(row=r, column=1, value=inst.get("aprendices_evaluaron") or 0)
        wsi.cell(row=r, column=2, value=len(inst.get("preguntas") or []))
        wsi.cell(row=r, column=3, value=inst.get("total_respuestas") or 0)
        wsi.cell(row=r, column=4, value=inst.get("total_siempre") or 0)
        wsi.cell(row=r, column=5, value=inst.get("total_algunas") or 0)
        wsi.cell(row=r, column=6, value=inst.get("total_nunca") or 0)
        pct = inst.get("porcentaje")
        wsi.cell(row=r, column=7, value=f"{pct} %" if pct is not None else "—")

        r += 2
        wsi.cell(row=r, column=1, value="OBSERVACIONES").font = Font(bold=True)
        r += 1
        obs = inst.get("observaciones") or []
        if not obs:
            wsi.cell(row=r, column=1, value="Ninguna")
        else:
            for o in obs:
                if o and str(o).strip():
                    wsi.cell(row=r, column=1, value=str(o).strip())
                    r += 1
        _auto_ancho(wsi)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pdf_reporte_instructor(data: dict) -> bytes:
    """
    Genera un PDF simple por instructor: foto (si hay path), nombre, correo,
    promedio general y promedio por pregunta (similar al consolidado Excel).
    data keys esperadas: nombre, apellido, correo, foto_path (opcional),
    promedio_general, preguntas (list of {orden, descripcion, nota, porcentaje, escala})
    """
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, HRFlowable
    )
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    import os

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, leftMargin=1.5*cm, rightMargin=1.5*cm, topMargin=1.5*cm, bottomMargin=1.5*cm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitleSENA", parent=styles["Heading1"], fontSize=16, textColor=colors.HexColor("#1a4a2e"), alignment=TA_CENTER, spaceAfter=6)
    sub_style = ParagraphStyle("SubSENA", parent=styles["Normal"], fontSize=10, textColor=colors.HexColor("#374151"), alignment=TA_CENTER, spaceAfter=12)
    h2 = ParagraphStyle("H2SENA", parent=styles["Heading2"], fontSize=12, textColor=colors.HexColor("#1a4a2e"), spaceBefore=10, spaceAfter=6)
    body = ParagraphStyle("BodySENA", parent=styles["Normal"], fontSize=9, leading=12)

    story = []
    story.append(Paragraph("SENA — Evaluación de Instructores", title_style))
    story.append(Paragraph("Reporte individual de resultados", sub_style))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#2d7a4f"), spaceAfter=12))

    nombre = f"{data.get('nombre') or ''} {data.get('apellido') or ''}".strip() or "Instructor"
    correo = data.get("correo") or "—"
    foto = data.get("foto_path")

    header_data = []
    if foto and os.path.isfile(foto):
        try:
            img = Image(foto, width=2.2*cm, height=2.2*cm)
            header_data.append([img, Paragraph(f"<b>{nombre}</b><br/>{correo}", body)])
        except Exception:
            header_data.append([Paragraph(f"<b>{nombre}</b><br/>{correo}", body)])
    else:
        header_data.append([Paragraph(f"<b>{nombre}</b><br/>{correo}", body)])

    t = Table(header_data, colWidths=[3*cm, 14*cm] if len(header_data[0]) == 2 else [17*cm])
    t.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(t)
    story.append(Spacer(1, 10))

    prom = data.get("promedio_general")
    if prom is not None:
        story.append(Paragraph(f"<b>Promedio general:</b> {prom:.2f} / 5.0  ({(prom/5)*100:.1f} %)", h2))
    else:
        story.append(Paragraph("<b>Promedio general:</b> Sin datos", h2))

    story.append(Paragraph("Detalle por pregunta", h2))
    preguntas = data.get("preguntas") or []
    if not preguntas:
        story.append(Paragraph("No hay respuestas registradas para este instructor.", body))
    else:
        rows = [["#", "Pregunta", "Nota", "%", "Escala"]]
        for p in preguntas:
            desc = (p.get("descripcion") or "")[:80]
            rows.append([
                str(p.get("orden") or ""),
                Paragraph(desc, body),
                f"{p.get('nota', 0):.2f}",
                f"{p.get('porcentaje', 0):.0f}%",
                p.get("escala") or "—",
            ])
        tabla = Table(rows, colWidths=[1*cm, 10*cm, 1.8*cm, 1.5*cm, 2.5*cm])
        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a4a2e")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d1d5db")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f4f6")]),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(tabla)

    story.append(Spacer(1, 16))
    story.append(Paragraph("Documento generado automáticamente por el sistema de Evaluación de Instructores SENA.", ParagraphStyle("Foot", parent=styles["Normal"], fontSize=7, textColor=colors.gray, alignment=TA_CENTER)))

    doc.build(story)
    return buf.getvalue()
