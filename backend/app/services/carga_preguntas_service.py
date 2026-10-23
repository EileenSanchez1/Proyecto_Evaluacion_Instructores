"""
Carga de preguntas desde Excel del formato oficial SENA
(Evaluación de instructores en aula).

Se toman únicamente filas que parecen preguntas:
- Columna con número de orden (ej. 1.1, 2.3) o entero
- Columna con el texto de la afirmación / pregunta

Formato esperado (flexible):
  | orden | descripcion |
  | 1.1   | El instructor es puntual... |
o el layout del formato SENA donde la primera columna es el código
y la segunda el texto.
"""
from __future__ import annotations

import io
import re
from typing import Any

from sqlmodel import Session, select

from app.models.pregunta import Pregunta


def _celda(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def parsear_xlsx_preguntas(contenido: bytes) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise RuntimeError(
            "openpyxl no está instalado. Ejecuta: pip install openpyxl"
        ) from e

    wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    ws = wb.active
    filas: list[dict[str, Any]] = []
    orden_auto = 1

    # Patrones de códigos de pregunta del formato SENA (1.1, 2.4, 3.2, etc.)
    codigo_re = re.compile(r"^(\d+)(?:\.(\d+))?$")

    for row in ws.iter_rows(values_only=True):
        if not row:
            continue
        celdas = [_celda(c) for c in row]
        if not any(celdas):
            continue

        # Saltar títulos / encabezados
        unidos = " ".join(celdas).lower()
        if any(
            x in unidos
            for x in (
                "servicio nacional",
                "evaluacion de instructores",
                "nombre instructor",
                "cordial saludo",
                "proposito de evaluar",
                "afirmaciones",
                "cumplimiento, conocimiento",
                "metodologia de formacion",
                "evaluacion aplicada",
                "clima organizacional",
            )
        ):
            # Si la fila es solo un título de sección (sin código de pregunta), omitir
            if not codigo_re.match(celdas[0].replace(" ", "")):
                continue

        # Buscar código en primera columna no vacía
        codigo = ""
        texto = ""
        for i, c in enumerate(celdas):
            if not c:
                continue
            limpio = c.replace(" ", "")
            if codigo_re.match(limpio) and not texto:
                codigo = limpio
                # texto = resto de la fila
                resto = [x for j, x in enumerate(celdas) if j != i and x]
                texto = " ".join(resto).strip()
                break

        if not codigo:
            # Intentar: primera col número, segunda texto
            if celdas[0] and len(celdas) > 1 and celdas[1]:
                if codigo_re.match(celdas[0].replace(" ", "")) or celdas[0].isdigit():
                    codigo = celdas[0].replace(" ", "")
                    texto = celdas[1]
            # O fila con solo texto largo (sin SI/NO headers)
            elif len(celdas[0]) > 25 and "si" not in celdas[0].lower()[:5]:
                texto = celdas[0]
                codigo = str(orden_auto)

        if not texto or len(texto) < 15:
            continue
        # Evitar columnas SI/NO sueltas
        if texto.upper() in ("SI", "NO", "SÍ"):
            continue

        # Orden numérico a partir del código (1.1 -> 11, 2.4 -> 24) o secuencial
        m = codigo_re.match(codigo.replace(" ", ""))
        if m:
            mayor = int(m.group(1))
            menor = int(m.group(2) or 0)
            orden = mayor * 10 + menor if menor else mayor
        else:
            orden = orden_auto

        filas.append({"orden": orden, "descripcion": texto, "codigo": codigo})
        orden_auto += 1

    wb.close()
    # Reordenar por orden
    filas.sort(key=lambda x: x["orden"])
    # Reasignar orden secuencial limpio 1..n
    for i, f in enumerate(filas, 1):
        f["orden"] = i
    return filas


def parsear_csv_preguntas(contenido: bytes) -> list[dict[str, Any]]:
    """
    Parsea CSV/TSV simple con columnas orden,descripcion o codigo,texto.
    Acepta separador coma o punto y coma o tab.
    """
    import csv

    texto = contenido.decode("utf-8-sig", errors="replace")
    # Detectar separador
    primera = texto.splitlines()[0] if texto.strip() else ""
    if "\t" in primera:
        delim = "\t"
    elif ";" in primera:
        delim = ";"
    else:
        delim = ","

    reader = csv.reader(io.StringIO(texto), delimiter=delim)
    filas: list[dict[str, Any]] = []
    orden_auto = 1
    codigo_re = re.compile(r"^(\d+)(?:\.(\d+))?$")
    header_seen = False

    for row in reader:
        if not row:
            continue
        celdas = [c.strip() for c in row]
        if not any(celdas):
            continue
        unidos = " ".join(celdas).lower()
        if any(
            x in unidos
            for x in ("orden", "descripcion", "pregunta", "codigo", "afirmacion", "texto")
        ) and not header_seen:
            header_seen = True
            continue
        header_seen = True

        codigo = ""
        texto_p = ""
        for i, c in enumerate(celdas):
            if not c:
                continue
            if not codigo and codigo_re.match(c.replace(" ", "")):
                codigo = c
                continue
            if len(c) >= 10:
                texto_p = c
                break
        if not texto_p and len(celdas) >= 2:
            texto_p = celdas[1] if len(celdas[1]) >= 5 else (celdas[-1] if celdas else "")
        if not texto_p or len(texto_p) < 10:
            continue
        m = codigo_re.match((codigo or "").replace(" ", ""))
        if m:
            mayor = int(m.group(1))
            menor = int(m.group(2) or 0)
            orden = mayor * 10 + menor if menor else mayor
        else:
            orden = orden_auto
        filas.append({"orden": orden, "descripcion": texto_p, "codigo": codigo})
        orden_auto += 1

    filas.sort(key=lambda x: x["orden"])
    for i, f in enumerate(filas, 1):
        f["orden"] = i
    return filas


def cargar_preguntas_desde_filas(
    session: Session,
    filas: list[dict[str, Any]],
    reemplazar: bool = False,
) -> dict[str, Any]:
    if reemplazar:
        existentes = session.exec(select(Pregunta)).all()
        for p in existentes:
            p.estado = False
            session.add(p)
        session.commit()

    creadas = 0
    actualizadas = 0
    omitidas = 0

    for item in filas:
        desc = (item.get("descripcion") or "").strip()
        orden = int(item.get("orden") or 0)
        if not desc:
            continue
        ya = session.exec(select(Pregunta).where(Pregunta.descripcion == desc)).first()
        if ya:
            if reemplazar:
                ya.estado = True
                ya.orden = orden
                session.add(ya)
                actualizadas += 1
            else:
                omitidas += 1
            continue
        session.add(Pregunta(descripcion=desc, orden=orden, estado=True))
        creadas += 1

    session.commit()
    return {
        "mensaje": f"Preguntas desde Excel: {creadas} creadas, {actualizadas} actualizadas, {omitidas} omitidas.",
        "creadas": creadas,
        "actualizadas": actualizadas,
        "omitidas": omitidas,
        "total_en_archivo": len(filas),
        "reemplazar": reemplazar,
    }
