"""
Carga masiva de aprendices desde CSV o Excel (.xlsx / .xls).

Formato oficial SENA — «Reporte de Aprendices»:

  (filas superiores pueden traer título y metadatos)
  Ficha de Caracterización: 3171690 - PRODUCCION DE MULTIMEDIA
  ...
  Tipo de Documento | Número de Documento | Nombre | Apellidos |
  Celular | Correo Electrónico | Estado

También se acepta el formato simple CSV:
  nombre,apellido,correo,numero_ficha

Cada fila:
  - Crea aprendiz + usuario ya asignado a la ficha
  - Genera contraseña temporal
  - Envía correo con usuario/contraseña (cartero)
"""

from __future__ import annotations

import csv
import io
import os
import re
import subprocess
import tempfile
from typing import Any

from sqlmodel import Session, select

from app.models.ficha import Ficha
from app.schemas.aprendiz import AprendizCreate
from app.services.aprendiz_service import AprendizService


def _normalizar_cabecera(nombre: str) -> str:
    n = (nombre or "").strip().lower()
    n = (
        n.replace(" ", "_")
        .replace("-", "_")
        .replace(".", "")
        .replace("á", "a")
        .replace("é", "e")
        .replace("í", "i")
        .replace("ó", "o")
        .replace("ú", "u")
        .replace("ñ", "n")
    )
    n = re.sub(r"[^a-z0-9_]", "", n)

    alias = {
        "email": "correo",
        "mail": "correo",
        "e_mail": "correo",
        "correo_electronico": "correo",
        "correoelectronico": "correo",
        "ficha": "numero_ficha",
        "n_ficha": "numero_ficha",
        "num_ficha": "numero_ficha",
        "numero_ficha": "numero_ficha",
        "nombres": "nombre",
        "nombre": "nombre",
        "apellidos": "apellido",
        "apellido": "apellido",
        "periodo": "id_periodo",
        "idperiodo": "id_periodo",
        "tipo_de_documento": "tipo_documento",
        "tipodedocumento": "tipo_documento",
        "numero_de_documento": "documento",
        "numerodedocumento": "documento",
        "numero_documento": "documento",
        "documento": "documento",
        "cedula": "documento",
        "celular": "celular",
        "telefono": "celular",
        "tel": "celular",
        "estado": "estado",
    }
    return alias.get(n, n)


def _extraer_numero_ficha_metadato(texto: str) -> str | None:
    """Busca 'Ficha de Caracterización: 3171690 - PROGRAMA' o similar."""
    if not texto:
        return None
    m = re.search(
        r"ficha\s*(?:de\s*caracterizaci[oó]n)?\s*[:\-]?\s*(\d{5,})",
        texto,
        re.IGNORECASE,
    )
    if m:
        return m.group(1)
    m = re.search(r"\b(\d{6,})\b", texto)
    return m.group(1) if m else None


def _filas_desde_matriz(
    rows: list[list[Any]], ficha_meta: str | None = None
) -> list[dict[str, str]]:
    """Convierte matriz de celdas (lista de filas) en filas normalizadas."""
    if not rows:
        return []

    if ficha_meta is None:
        for row in rows[:15]:
            for cell in row:
                if cell is None:
                    continue
                num = _extraer_numero_ficha_metadato(str(cell))
                if num:
                    ficha_meta = num
                    break
            if ficha_meta:
                break

    idx_header = 0
    for i, row in enumerate(rows[:40]):
        textos = [str(c).strip().lower() if c is not None else "" for c in row]
        joined = " ".join(textos)
        if (
            ("nombre" in joined and ("apellido" in joined or "correo" in joined))
            or "tipo de documento" in joined
            or "correo electrónico" in joined
            or "correo electronico" in joined
        ):
            idx_header = i
            break

    headers = [
        _normalizar_cabecera(str(h) if h is not None else "")
        for h in rows[idx_header]
    ]

    filas: list[dict[str, str]] = []
    for row in rows[idx_header + 1 :]:
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        normalizada: dict[str, str] = {}
        for i, cell in enumerate(row):
            if i >= len(headers):
                break
            key = headers[i]
            if not key:
                continue
            normalizada[key] = str(cell).strip() if cell is not None else ""
        if ficha_meta and not normalizada.get("numero_ficha"):
            normalizada["numero_ficha"] = ficha_meta
        filas.append(normalizada)
    return filas


def parsear_csv(contenido: bytes | str) -> list[dict[str, str]]:
    if isinstance(contenido, bytes):
        if contenido.startswith(b"\xef\xbb\xbf"):
            contenido = contenido[3:]
        texto = contenido.decode("utf-8", errors="replace")
    else:
        texto = contenido

    lineas = texto.splitlines()
    ficha_meta: str | None = None
    for linea in lineas[:15]:
        num = _extraer_numero_ficha_metadato(linea)
        if num:
            ficha_meta = num
            break

    idx_header = 0
    for i, linea in enumerate(lineas[:30]):
        low = linea.lower()
        if (
            ("nombre" in low and ("apellido" in low or "correo" in low))
            or "tipo de documento" in low
            or "correo electrónico" in low
            or "correo electronico" in low
        ):
            idx_header = i
            break

    texto_util = "\n".join(lineas[idx_header:])
    primera = texto_util.splitlines()[0] if texto_util.strip() else ""
    delimiter = ";" if primera.count(";") > primera.count(",") else ","

    reader = csv.DictReader(io.StringIO(texto_util), delimiter=delimiter)
    filas: list[dict[str, str]] = []
    for row in reader:
        normalizada: dict[str, str] = {}
        for k, v in row.items():
            if k is None:
                continue
            key = _normalizar_cabecera(k)
            normalizada[key] = (v or "").strip()
        if ficha_meta and not normalizada.get("numero_ficha"):
            normalizada["numero_ficha"] = ficha_meta
        filas.append(normalizada)
    return filas


def parsear_xlsx_bytes(contenido: bytes) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return _filas_desde_matriz(rows)


def _es_html(contenido: bytes) -> bool:
    head = contenido[:2000].lstrip().lower()
    return head.startswith(b"<html") or head.startswith(b"<!doctype") or b"<table" in head


def parsear_html_xls(contenido: bytes) -> list[dict[str, str]]:
    """Muchos reportes SENA se exportan como HTML con extensión .xls."""
    try:
        import pandas as pd
    except ImportError as e:
        raise RuntimeError("pandas no está instalado para leer HTML/.xls") from e

    texto = contenido.decode("utf-8", errors="replace")
    # Metadato de ficha en el HTML
    ficha_meta = _extraer_numero_ficha_metadato(texto)

    tablas = pd.read_html(io.StringIO(texto))
    if not tablas:
        raise ValueError("No se encontraron tablas en el archivo HTML/.xls.")

    # Elegir la tabla con más filas / que tenga columnas de nombre/correo
    mejor = None
    for t in tablas:
        cols = " ".join(str(c).lower() for c in t.columns)
        if "nombre" in cols or "correo" in cols or "documento" in cols:
            mejor = t
            break
    if mejor is None:
        mejor = max(tablas, key=lambda t: len(t))

    # A veces la primera fila de datos es la cabecera real
    rows: list[list[Any]] = [list(mejor.columns)]
    for _, series in mejor.iterrows():
        rows.append([None if (isinstance(v, float) and str(v) == "nan") else v for v in series.tolist()])
    return _filas_desde_matriz(rows, ficha_meta=ficha_meta)


def _convertir_xls_a_xlsx(contenido: bytes) -> bytes:
    """Convierte .xls antiguo a .xlsx con LibreOffice (si está disponible)."""
    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, "entrada.xls")
        with open(src, "wb") as f:
            f.write(contenido)
        # soffice --headless --convert-to xlsx
        cmd = [
            "soffice",
            "--headless",
            "--norestore",
            "--nolockcheck",
            "--convert-to",
            "xlsx",
            "--outdir",
            tmp,
            src,
        ]
        try:
            subprocess.run(
                cmd,
                check=True,
                capture_output=True,
                timeout=60,
            )
        except FileNotFoundError:
            raise RuntimeError(
                "No se pudo leer el .xls. Guarda el archivo como .xlsx desde Excel "
                "o instala LibreOffice en el servidor."
            )
        except subprocess.CalledProcessError as e:
            err = (e.stderr or b"").decode("utf-8", errors="replace")[:300]
            raise RuntimeError(
                f"No se pudo convertir el .xls a .xlsx. "
                f"Guarda el archivo como .xlsx e inténtalo de nuevo. Detalle: {err}"
            ) from e
        except subprocess.TimeoutExpired as e:
            raise RuntimeError(
                "Tiempo agotado al convertir el .xls. Guarda el archivo como .xlsx."
            ) from e

        out = os.path.join(tmp, "entrada.xlsx")
        if not os.path.isfile(out):
            # A veces el nombre de salida varía
            candidatos = [p for p in os.listdir(tmp) if p.endswith(".xlsx")]
            if not candidatos:
                raise RuntimeError(
                    "La conversión a .xlsx no generó archivo. "
                    "Abre el reporte en Excel y guárdalo como .xlsx."
                )
            out = os.path.join(tmp, candidatos[0])
        with open(out, "rb") as f:
            return f.read()


def parsear_excel(contenido: bytes, filename: str = "") -> list[dict[str, str]]:
    """
    Lee .xlsx (openpyxl), .xls HTML (pandas) o .xls binario (LibreOffice → xlsx).
    """
    nombre = (filename or "").lower()

    # 1) HTML disfrazado de .xls (muy común en reportes SENA)
    if _es_html(contenido):
        return parsear_html_xls(contenido)

    # 2) .xlsx real
    if nombre.endswith(".xlsx") or contenido[:2] == b"PK":
        try:
            return parsear_xlsx_bytes(contenido)
        except Exception as e:
            # Si falla y parece HTML mal detectado
            if _es_html(contenido):
                return parsear_html_xls(contenido)
            raise RuntimeError(
                f"No se pudo leer el Excel (.xlsx): {type(e).__name__}: {e}"
            ) from e

    # 3) .xls antiguo (BIFF) → convertir con LibreOffice
    if nombre.endswith(".xls") or nombre.endswith(".xlsb"):
        try:
            xlsx_bytes = _convertir_xls_a_xlsx(contenido)
            return parsear_xlsx_bytes(xlsx_bytes)
        except Exception as e:
            # Último intento: HTML
            if _es_html(contenido):
                return parsear_html_xls(contenido)
            raise RuntimeError(
                "No se pudo leer el archivo .xls. "
                "Ábrelo en Excel y guárdalo como «Libro de Excel (.xlsx)» e inténtalo de nuevo. "
                f"Detalle: {type(e).__name__}: {e}"
            ) from e

    # 4) Intento genérico openpyxl
    try:
        return parsear_xlsx_bytes(contenido)
    except Exception:
        if _es_html(contenido):
            return parsear_html_xls(contenido)
        raise


def parsear_archivo(contenido: bytes, filename: str) -> list[dict[str, str]]:
    nombre = (filename or "").lower()
    if (
        nombre.endswith(".xlsx")
        or nombre.endswith(".xls")
        or nombre.endswith(".xlsb")
    ):
        return parsear_excel(contenido, filename)
    return parsear_csv(contenido)


def cargar_desde_filas(
    session: Session,
    filas: list[dict[str, str]],
    id_periodo_default: int | None = None,
    enviar_correo: bool = True,
    id_ficha_forzada: int | None = None,
) -> dict[str, Any]:
    """
    Procesa filas ya parseadas.
    Si id_ficha_forzada se indica, todos los aprendices van a esa ficha
    (útil al subir el reporte dentro del detalle de una ficha).
    """
    fichas = {f.numero_ficha: f for f in session.exec(select(Ficha)).all()}
    ficha_fija = session.get(Ficha, id_ficha_forzada) if id_ficha_forzada else None

    creados = []
    errores = []
    omitidos = []

    for i, fila in enumerate(filas, start=2):
        nombre = (fila.get("nombre") or "").strip()
        apellido = (fila.get("apellido") or "").strip()
        correo = (fila.get("correo") or "").strip().lower()
        numero_ficha = (
            fila.get("numero_ficha")
            or fila.get("id_ficha")
            or ""
        ).strip()
        id_periodo_raw = (fila.get("id_periodo") or "").strip()

        if not nombre and not apellido and not correo:
            continue

        if not nombre or not apellido or not correo:
            errores.append({
                "fila": i,
                "correo": correo or "(sin correo)",
                "error": "Faltan campos obligatorios (Nombre, Apellidos, Correo Electrónico).",
            })
            continue

        if ficha_fija:
            ficha = ficha_fija
        else:
            if not numero_ficha:
                errores.append({
                    "fila": i,
                    "correo": correo,
                    "error": (
                        "No se encontró número de ficha. "
                        "El Excel debe traer «Ficha de Caracterización: NNNNNN» "
                        "o la columna numero_ficha."
                    ),
                })
                continue
            ficha = fichas.get(numero_ficha)
            if not ficha:
                try:
                    id_f = int(numero_ficha)
                    ficha = session.get(Ficha, id_f)
                except ValueError:
                    ficha = None
            if not ficha:
                errores.append({
                    "fila": i,
                    "correo": correo,
                    "error": (
                        f"Ficha '{numero_ficha}' no existe. "
                        "Créala antes de cargar aprendices."
                    ),
                })
                continue

        id_periodo = id_periodo_default
        if id_periodo_raw:
            try:
                id_periodo = int(id_periodo_raw)
            except ValueError:
                errores.append({
                    "fila": i,
                    "correo": correo,
                    "error": f"id_periodo inválido: {id_periodo_raw}",
                })
                continue

        try:
            data = AprendizCreate(
                nombre=nombre,
                apellido=apellido,
                correo=correo,
                id_ficha=ficha.id_ficha,
                id_periodo=id_periodo,
                enviar_correo=enviar_correo,
            )
            aprendiz = AprendizService.crear(session, data)
            creados.append({
                "fila": i,
                "correo": correo,
                "id_aprendiz": aprendiz.id_aprendiz,
                "numero_ficha": ficha.numero_ficha,
            })
        except ValueError as e:
            msg = str(e)
            if "Ya existe" in msg:
                omitidos.append({
                    "fila": i,
                    "correo": correo,
                    "motivo": msg,
                })
            else:
                errores.append({
                    "fila": i,
                    "correo": correo,
                    "error": msg,
                })
        except Exception as e:
            errores.append({
                "fila": i,
                "correo": correo,
                "error": f"{type(e).__name__}: {e}",
            })

    return {
        "total_filas": len(filas),
        "creados": len(creados),
        "omitidos": len(omitidos),
        "errores": len(errores),
        "detalle_creados": creados,
        "detalle_omitidos": omitidos,
        "detalle_errores": errores,
        "mensaje": (
            f"Proceso terminado: {len(creados)} creados, "
            f"{len(omitidos)} omitidos (ya existían), "
            f"{len(errores)} con error. "
            f"{'Se enviaron correos con usuario y contraseña.' if enviar_correo else 'No se enviaron correos.'}"
        ),
    }
